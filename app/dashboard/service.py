"""All dashboard logic: reads for the 3 dashboards, writes for the supervisor forms, audit log."""
import csv
import functools
import io
import re
import threading
import time
from datetime import date, datetime, time as dtime, timedelta
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from mysql.connector.errors import IntegrityError
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.database import get_cursor
from app.errors.exceptions import Conflict, InvalidRequest, NotFound

# Tiny TTL cache for expensive, read-only dashboard queries. Several users and
# the 10-30s browser polling all ask for the same data; compute it once.
_TTL_CACHE = {}
_TTL_LOCK = threading.Lock()


def ttl_cached(seconds):
    def decorator(fn):
        @functools.wraps(fn)
        def wrapper(*args, **kwargs):
            key = (fn.__name__, args, tuple(sorted(kwargs.items())))
            now = time.monotonic()
            with _TTL_LOCK:
                hit = _TTL_CACHE.get(key)
                if hit and now - hit[0] < seconds:
                    return hit[1]
            value = fn(*args, **kwargs)
            with _TTL_LOCK:
                _TTL_CACHE[key] = (time.monotonic(), value)
                if len(_TTL_CACHE) > 200:      # keep it bounded
                    _TTL_CACHE.clear()
            return value
        return wrapper
    return decorator


def clear_dashboard_cache():
    with _TTL_LOCK:
        _TTL_CACHE.clear()


DAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
DATA_ENTRY_ACTIONS = {"Shift Entry", "Downtime Logged", "Downtime Ended"}


# ---------------------------------------------------------------- helpers
def rows(sql, params=()):
    with get_cursor() as (_, cur):
        cur.execute(sql, params)
        return cur.fetchall()


def one(sql, params=()):
    with get_cursor() as (_, cur):
        cur.execute(sql, params)
        return cur.fetchone()


def pct(num, den):
    return round(float(num) / float(den) * 100, 1) if den else 0.0


def fmt_minutes(m):
    m = int(m or 0)
    return f"{m // 60}h {m % 60}m" if m >= 60 else f"{m}m"


def period(range_name):
    """Return (start, end, previous_start, previous_end) as dates."""
    today = date.today()
    if range_name in {"weekly", "week"}:
        start = today - timedelta(days=today.weekday())
        previous_end = start - timedelta(days=1)
        return start, today, previous_end - timedelta(days=6), previous_end
    if range_name in {"monthly", "month"}:
        start = today.replace(day=1)
        previous_end = start - timedelta(days=1)
        previous_start = previous_end.replace(day=1)
        return start, today, previous_start, previous_end
    return today, today, today - timedelta(days=1), today - timedelta(days=1)


def log_audit(user, action, target="", details=""):
    """Call this from anywhere (e.g. your login route: log_audit(user, "Login"))."""
    name = (user or {}).get("full_name") or "System"
    uid = (user or {}).get("user_id")
    with get_cursor() as (_, cur):
        cur.execute(
            "INSERT INTO audit_logs (user_id, user_name, action, target, details) VALUES (%s,%s,%s,%s,%s)",
            (uid, name, action, str(target)[:190], str(details)[:500]),
        )


def _pick(body, *names, default=None):
    for n in names:
        if n in body and body[n] not in (None, ""):
            return body[n]
    return default


def _int(value, label, minimum=0):
    try:
        n = int(float(value))
    except (TypeError, ValueError):
        raise InvalidRequest(f"{label} must be a number")
    if n < minimum:
        raise InvalidRequest(f"{label} cannot be below {minimum}")
    return n


def _find(table, id_col, code_col, name_col, value, label):
    """Resolve an id, code, name, or 'CODE Name' string from the UI to a row."""
    if value in (None, ""):
        raise InvalidRequest(f"{label} is required")
    text = str(value).strip()
    row = None
    if text.isdigit():
        row = one(f"SELECT * FROM {table} WHERE {id_col} = %s", (int(text),))
    if not row:
        row = one(
            f"SELECT * FROM {table} WHERE {code_col} = %s OR {name_col} = %s "
            f"OR CONCAT({code_col}, ' ', {name_col}) = %s",
            (text, text, text),
        )
    if not row:
        row = one(f"SELECT * FROM {table} WHERE {code_col} = %s", (text.split(" ")[0],))
    if not row:
        raise NotFound(f"{label} '{text}' was not found")
    return row


def find_product(v): return _find("products", "product_id", "product_code", "product_name", v, "Product")
def find_machine(v): return _find("machines", "machine_id", "machine_code", "machine_name", v, "Machine")
def find_shift(v):
    if v in (None, ""):
        return None
    text = str(v).strip()
    row = one("SELECT * FROM shifts WHERE shift_id = %s OR shift_name = %s", (int(text) if text.isdigit() else -1, text))
    if not row:  # UI strings look like "A – Day"; match on the first word too
        row = one("SELECT * FROM shifts WHERE shift_name LIKE %s", (text.split(" ")[0] + "%",))
    if not row:
        raise NotFound(f"Shift '{text}' was not found")
    return row


def _to_datetime(value, base=None):
    """Accepts 'HH:MM', 'YYYY-MM-DD HH:MM', or ISO strings."""
    if value in (None, ""):
        return None
    text = str(value).strip().replace("T", " ")
    base = base or date.today()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.combine(base, datetime.strptime(text, fmt).time())
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            pass
    raise InvalidRequest(f"Could not read time '{value}'. Use HH:MM.")


# ---------------------------------------------------------------- supervisor
DEFAULT_SHIFTS = [("A - Day", "06:00:00", "14:00:00"), ("B - Evening", "14:00:00", "22:00:00"), ("C - Night", "22:00:00", "06:00:00")]

EXTRA_TABLES = [
    """CREATE TABLE IF NOT EXISTS audit_logs (
         audit_id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
         user_id BIGINT UNSIGNED NULL,
         user_name VARCHAR(120) NOT NULL DEFAULT 'System',
         action VARCHAR(80) NOT NULL,
         target VARCHAR(190) NOT NULL DEFAULT '',
         details VARCHAR(500) NOT NULL DEFAULT '',
         created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
         INDEX idx_audit_created (created_at)
       ) ENGINE=InnoDB""",
    """CREATE TABLE IF NOT EXISTS defect_reasons (
         defect_id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
         production_id BIGINT UNSIGNED NOT NULL,
         cause VARCHAR(120) NOT NULL,
         quantity INT UNSIGNED NOT NULL DEFAULT 0,
         FOREIGN KEY (production_id) REFERENCES production_records(production_id) ON DELETE CASCADE
       ) ENGINE=InnoDB""",
    """CREATE TABLE IF NOT EXISTS target_progress_reports (
         report_id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
         target_id BIGINT UNSIGNED NOT NULL,
         supervisor_id BIGINT UNSIGNED NOT NULL,
         manager_id BIGINT UNSIGNED NOT NULL,
         actual_quantity BIGINT UNSIGNED NOT NULL,
         reason TEXT NOT NULL,
         created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
         reviewed_at TIMESTAMP NULL DEFAULT NULL,
         INDEX idx_target_report_manager (manager_id, created_at),
         FOREIGN KEY (target_id) REFERENCES production_targets(target_id),
         FOREIGN KEY (supervisor_id) REFERENCES users(user_id),
         FOREIGN KEY (manager_id) REFERENCES users(user_id)
       ) ENGINE=InnoDB""",
]
_schema_ready = False


_default_shifts_ready = False


def ensure_default_shifts():
    """Ensure A/B/C shifts exist without overwriting existing shift definitions.
    Runs once per process; it used to run on every supervisor dashboard load."""
    global _default_shifts_ready
    if _default_shifts_ready:
        return
    with get_cursor() as (_, cur):
        for name, st, en in DEFAULT_SHIFTS:
            cur.execute(
                "INSERT IGNORE INTO shifts (shift_name, start_time, end_time) VALUES (%s,%s,%s)",
                (name, st, en),
            )
    _default_shifts_ready = True


def ensure_schema():
    """Create dashboard tables, migrate supervisor assignments, and ensure default shifts."""
    global _schema_ready
    if _schema_ready:
        return
    with get_cursor() as (_, cur):
        for ddl in EXTRA_TABLES:
            cur.execute(ddl)
        cur.execute("SHOW COLUMNS FROM users LIKE 'shift_id'")
        if not cur.fetchone():
            cur.execute(
                """ALTER TABLE users
                   ADD COLUMN shift_id BIGINT UNSIGNED NULL,
                   ADD INDEX idx_users_shift (shift_id),
                   ADD CONSTRAINT fk_users_shift
                     FOREIGN KEY (shift_id) REFERENCES shifts(shift_id) ON DELETE SET NULL"""
            )
    ensure_default_shifts()
    _schema_ready = True


def supervisor_dashboard(user=None):
    ensure_default_shifts()
    today = date.today()
    products = rows("SELECT product_code, product_name FROM products WHERE is_active = 1 ORDER BY product_code")
    machines = rows("SELECT machine_name FROM machines ORDER BY machine_name")
    shifts = rows("SELECT shift_name FROM shifts WHERE is_active = 1 ORDER BY start_time")
    target = one("SELECT COALESCE(SUM(target_quantity),0) AS t FROM production_targets WHERE target_date = %s", (today,))["t"]
    produced = one("SELECT COALESCE(SUM(quantity_produced),0) AS p FROM production_records WHERE production_date = %s", (today,))["p"]

    targets = [{"product": t["product"], "machine": t["machine"], "shift": t["shift"], "quantity": int(t["quantity"])} for t in rows(
        """SELECT CONCAT(p.product_code, ' ', p.product_name) AS product, m.machine_name AS machine, s.shift_name AS shift,
                  t.target_quantity AS quantity
           FROM production_targets t JOIN products p ON p.product_id = t.product_id
           LEFT JOIN machines m ON m.machine_id = t.machine_id LEFT JOIN shifts s ON s.shift_id = t.shift_id
           WHERE t.target_date = %s""", (today,))]

    downtime = rows(
        """SELECT m.machine_name, d.reason, d.started_at, d.ended_at
           FROM downtime_events d JOIN machines m ON m.machine_id = d.machine_id
           WHERE DATE(d.started_at) = %s ORDER BY d.started_at DESC LIMIT 20""",
        (today,),
    )
    now = datetime.now()
    down_out = [
        {
            "machine": d["machine_name"],
            "reason": d["reason"],
            "start": d["started_at"].strftime("%H:%M"),
            "end": d["ended_at"].strftime("%H:%M") if d["ended_at"] else "Ongoing",
            "duration": fmt_minutes(((d["ended_at"] or now) - d["started_at"]).total_seconds() / 60),
        }
        for d in downtime
    ]

    first = today - timedelta(days=6)
    per_day = {
        r["d"]: r
        for r in rows(
            """SELECT DATE(started_at) AS d, COUNT(*) AS events,
                      COALESCE(SUM(TIMESTAMPDIFF(MINUTE, started_at, COALESCE(ended_at, %s))),0) AS minutes
               FROM downtime_events WHERE DATE(started_at) BETWEEN %s AND %s GROUP BY DATE(started_at)""",
            (now, first, today),
        )
    }
    trend = []
    for i in range(7):
        d = first + timedelta(days=i)
        r = per_day.get(d)
        trend.append({"day": DAY_NAMES[d.weekday()], "events": int(r["events"]) if r else 0,
                      "minutes": int(r["minutes"]) if r else 0})

    assigned_shift = None
    if user and user.get("user_id"):
        assigned_shift = one(
            """SELECT s.shift_id, s.shift_name, s.start_time, s.end_time
               FROM users u LEFT JOIN shifts s ON s.shift_id = u.shift_id
               WHERE u.user_id = %s""",
            (user["user_id"],),
        )
        if assigned_shift and assigned_shift["shift_id"] is None:
            assigned_shift = None

    return {
        "products": [f"{p['product_code']} {p['product_name']}" for p in products],
        "machines": [m["machine_name"] for m in machines],
        "shifts": [s["shift_name"] for s in shifts],
        "assigned_shift": assigned_shift,
        "target": int(target),
        "targets": targets,
        "produced": int(produced),
        "downtime": down_out,
        "trend": trend,
    }


def submit_shift_entry(body, user):
    product = find_product(_pick(body, "product", "product_id", "product_code"))
    machine = find_machine(_pick(body, "machine", "machine_id", "machine_code"))
    shift = find_shift(_pick(body, "shift", "shift_id", "shift_name"))
    if user.get("role") == "SUPERVISOR" and user.get("shift_id"):
        assigned_shift = find_shift(user["shift_id"])
        if not assigned_shift["is_active"]:
            raise InvalidRequest("Your assigned shift is inactive. Contact a manager to update it.")
        if shift and shift["shift_id"] != assigned_shift["shift_id"]:
            raise InvalidRequest(
                f"Production entries must use your assigned shift ({assigned_shift['shift_name']}). "
                "Contact a manager to request a shift change."
            )
        shift = assigned_shift
    produced = _int(_pick(body, "produced", "quantity_produced", "total", "actual", "output", "quantity", default=0), "Produced quantity")
    defective = _int(_pick(body, "defects", "defective", "quantity_defective", "scrap", default=0), "Defective quantity")
    if defective > produced:
        raise InvalidRequest("Defective quantity cannot be more than produced quantity")
    if body.get("good") not in (None, "") and _int(body["good"], "Good quantity") + defective != produced:
        raise InvalidRequest("Good + defective quantity must equal total manufactured")
    run = _pick(body, "run_minutes", "runMinutes", "runtime")
    planned = _pick(body, "planned_minutes", "plannedMinutes", "planned")
    run = _int(run, "Run minutes") if run is not None else None
    planned = _int(planned, "Planned minutes") if planned is not None else None
    when = _pick(body, "date", "production_date")
    prod_date = _to_datetime(f"{when} 00:00").date() if when else date.today()

    reasons = body.get("defect_reasons") or []  # optional: [{"cause": "Seal fail", "quantity": 4}]
    if sum(_int(r.get("quantity", 0), "Defect reason quantity") for r in reasons) > defective:
        raise InvalidRequest("Defect reasons add up to more than the defective quantity")

    with get_cursor() as (_, cur):
        cur.execute(
            """INSERT INTO production_records
               (product_id, machine_id, shift_id, supervisor_id, production_date,
                quantity_produced, quantity_defective, run_minutes, planned_minutes, notes)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            (product["product_id"], machine["machine_id"], shift["shift_id"] if shift else None,
             user["user_id"], prod_date, produced, defective, run, planned, _pick(body, "notes", "note")),
        )
        production_id = cur.lastrowid
        for r in reasons:
            cur.execute("INSERT INTO defect_reasons (production_id, cause, quantity) VALUES (%s,%s,%s)",
                        (production_id, str(r.get("cause", "Other"))[:120], _int(r.get("quantity", 0), "Defect reason quantity")))
    log_audit(user, "Shift Entry", f"{product['product_code']} on {machine['machine_name']}",
              f"{produced} produced, {defective} defective")
    return {"production_id": production_id, "ok": True}


def log_downtime(body, user):
    machine = find_machine(_pick(body, "machine", "machine_id", "machine_code"))
    reason = _pick(body, "reason", "cause")
    if not reason:
        raise InvalidRequest("Reason is required")
    start = _to_datetime(_pick(body, "start", "started_at", "from"))
    if start is None:
        start = datetime.now()
    end = _to_datetime(_pick(body, "end", "ended_at", "to"), base=start.date())
    if end and end < start:
        raise InvalidRequest("End time cannot be before start time")
    with get_cursor() as (_, cur):
        cur.execute(
            "INSERT INTO downtime_events (machine_id, started_at, ended_at, reason, notes, reported_by) VALUES (%s,%s,%s,%s,%s,%s)",
            (machine["machine_id"], start, end, str(reason)[:255], _pick(body, "notes", "note"), user["user_id"]),
        )
        downtime_id = cur.lastrowid
        if end is None:  # still down
            cur.execute("UPDATE machines SET status = 'DOWN' WHERE machine_id = %s", (machine["machine_id"],))
    log_audit(user, "Downtime Logged", machine["machine_name"], f"{reason}" + (" (ongoing)" if end is None else ""))
    return {"downtime_id": downtime_id, "ok": True}


def end_downtime(downtime_id, user):
    with get_cursor() as (_, cur):
        cur.execute("SELECT machine_id, ended_at FROM downtime_events WHERE downtime_id = %s", (downtime_id,))
        d = cur.fetchone()
        if not d:
            raise NotFound("Downtime event not found")
        if d["ended_at"]:
            raise Conflict("This downtime has already ended")
        cur.execute("UPDATE downtime_events SET ended_at = %s WHERE downtime_id = %s", (datetime.now(), downtime_id))
        cur.execute("UPDATE machines SET status = 'RUNNING' WHERE machine_id = %s AND status = 'DOWN'", (d["machine_id"],))
    log_audit(user, "Downtime Ended", f"Event #{downtime_id}")
    return {"ok": True}


# ---------------------------------------------------------------- manager

def _per_machine(start, end):
    """Per-machine production totals plus OEE parts for a period."""
    base = rows(
        """SELECT m.machine_id, m.machine_code, m.machine_name, m.location, m.status,
                  COALESCE(SUM(r.quantity_produced),0) AS produced, COALESCE(SUM(r.quantity_defective),0) AS defective,
                  COALESCE(SUM(r.run_minutes),0) AS run, COALESCE(SUM(r.planned_minutes),0) AS planned
           FROM machines m LEFT JOIN production_records r ON r.machine_id = m.machine_id AND r.production_date BETWEEN %s AND %s
           GROUP BY m.machine_id ORDER BY m.machine_code""", (start, end))
    targets = {x["machine_id"]: float(x["t"]) for x in rows(
        "SELECT machine_id, SUM(target_quantity) AS t FROM production_targets WHERE machine_id IS NOT NULL AND target_date BETWEEN %s AND %s GROUP BY machine_id",
        (start, end))}
    out = []
    for x in base:
        o = _oee(x["produced"], x["defective"], x["run"], x["planned"], targets.get(x["machine_id"], 0))
        out.append({**x, "target": targets.get(x["machine_id"], 0), **o})
    return out


@ttl_cached(10)
def compute_alerts(start, end):
    now = datetime.now()
    alerts = []
    for d in rows(
        """SELECT d.downtime_id, m.machine_name, d.started_at, d.ended_at
           FROM downtime_events d JOIN machines m ON m.machine_id = d.machine_id
           WHERE DATE(d.started_at) BETWEEN %s AND %s""", (start, end)):
        mins = ((d["ended_at"] or now) - d["started_at"]).total_seconds() / 60
        if mins > 60:
            alerts.append({"id": f"downtime:{d['downtime_id']}", "level": "high", "machine": d["machine_name"], "text": "Extended Downtime (> 60m)",
                           "duration": fmt_minutes(mins), "action": "Action Required"})
    for x in _per_machine(start, end):
        rate = pct(x["defective"], x["produced"])
        if rate > 5:
            alerts.append({"level": "medium", "machine": x["machine_name"], "text": "High Defect Rate (> 5%)",
                           "duration": f"{rate}%", "action": "Review quality log"})
        if x["planned"] and x["value"] < 80:
            alerts.append({"level": "low", "machine": x["machine_name"], "text": "OEE Below 80%",
                           "duration": f"{x['value']}%", "action": "Monitor closely"})
    alerts.sort(key=lambda a: {"high": 0, "medium": 1, "low": 2}[a["level"]])
    return alerts


def list_notifications(user):
    today = date.today()
    notifications = [
        {
            "id": alert.get("id") or f"database-alert:{alert['machine']}:{alert['text']}",
            "level": alert["level"],
            "text": f"{alert['machine']}: {alert['text']}",
            "time": f"{alert['duration']} · {alert['action']}",
        }
        for alert in compute_alerts(today, today)
    ]
    if user["role"] in {"MANAGER", "ADMIN"}:
        reports = rows(
            """SELECT r.report_id, r.actual_quantity, r.reason, r.created_at,
                      t.target_quantity, p.product_code, p.product_name,
                      sender.full_name AS supervisor_name
               FROM target_progress_reports r
               JOIN production_targets t ON t.target_id = r.target_id
               JOIN products p ON p.product_id = t.product_id
               JOIN users sender ON sender.user_id = r.supervisor_id
               WHERE r.manager_id = %s AND r.reviewed_at IS NULL
               ORDER BY r.created_at DESC
               LIMIT 50""",
            (user["user_id"],),
        )
        notifications.extend(
            {
                "id": f"target-report:{report['report_id']}",
                "level": "high",
                "text": (
                    f"{report['supervisor_name']} reported an at-risk target: "
                    f"{report['product_code']} {report['product_name']} "
                    f"({int(report['actual_quantity']):,}/{int(report['target_quantity']):,})"
                ),
                "time": report["created_at"].strftime("%d %b %H:%M"),
            }
            for report in reports
        )
    return notifications


FLOOR_STATUS = {"RUNNING": "running", "IDLE": "idle", "DOWN": "down", "MAINTENANCE": "maint"}


def machine_floor(start, end):
    open_reason = {x["machine_id"]: x["reason"] for x in rows(
        "SELECT machine_id, reason FROM downtime_events WHERE ended_at IS NULL ORDER BY started_at")}
    out = []
    for x in _per_machine(start, end):
        status = FLOOR_STATUS.get(x["status"], "idle")
        out.append({"name": x["machine_code"], "fullName": x["machine_name"], "status": status,
                    "oee": x["value"] if status != "down" else 0,
                    "reason": open_reason.get(x["machine_id"], "Scheduled service" if status == "maint" else "")})
    return out


def analytics(range_name="weekly"):
    start, end, _, _ = period(range_name)
    machines = _per_machine(start, end)
    lines = {}
    for m in machines:
        key = m["location"] or "Unassigned"
        L = lines.setdefault(key, {"line": key, "machines": 0, "running": 0, "produced": 0, "defective": 0, "run": 0, "planned": 0, "target": 0})
        L["machines"] += 1
        L["running"] += 1 if m["status"] == "RUNNING" else 0
        for k in ("produced", "defective", "run", "planned", "target"):
            L[k] += float(m[k])
    line_rows = []
    for L in lines.values():
        o = _oee(L["produced"], L["defective"], L["run"], L["planned"], L["target"])
        line_rows.append({"line": L["line"], "machines": L["machines"], "running": L["running"], "produced": int(L["produced"]),
                          "defectRate": pct(L["defective"], L["produced"]), "availability": o["availability"], "oee": o["value"]})
    return {
        "machines": [{"machine": m["machine_name"], "code": m["machine_code"], "line": m["location"] or "Unassigned",
                      "status": m["status"], "produced": int(m["produced"]), "defectRate": pct(m["defective"], m["produced"]),
                      "availability": m["availability"], "performance": m["performance"], "quality": m["quality"], "oee": m["value"]}
                     for m in machines],
        "lines": line_rows,
    }


def quality_summary():
    start, end, _, _ = period("weekly")
    by_machine = [{"machine": m["machine_name"], "produced": int(m["produced"]), "defective": int(m["defective"]),
                   "rate": pct(m["defective"], m["produced"])} for m in _per_machine(start, end)]
    recent = rows(
        """SELECT r.production_date, p.product_code, m.machine_name, r.quantity_produced, r.quantity_defective, u.full_name
           FROM production_records r JOIN products p ON p.product_id = r.product_id JOIN machines m ON m.machine_id = r.machine_id
           LEFT JOIN users u ON u.user_id = r.supervisor_id ORDER BY r.production_id DESC LIMIT 15""")
    return {"byMachine": by_machine,
            "recent": [{"date": str(x["production_date"]), "product": x["product_code"], "machine": x["machine_name"],
                        "produced": int(x["quantity_produced"]), "defective": int(x["quantity_defective"]),
                        "rate": pct(x["quantity_defective"], x["quantity_produced"]), "by": x["full_name"] or "-"} for x in recent]}


def set_machine_status(machine_id, status, user):
    status = str(status).upper()
    if status not in {"RUNNING", "IDLE", "DOWN", "MAINTENANCE"}:
        raise InvalidRequest("Status must be RUNNING, IDLE, DOWN or MAINTENANCE")
    with get_cursor() as (_, cur):
        cur.execute("UPDATE machines SET status = %s WHERE machine_id = %s", (status, machine_id))
        cur.execute("SELECT machine_name FROM machines WHERE machine_id = %s", (machine_id,))
        row = cur.fetchone()
    if not row:
        raise NotFound("Machine not found")
    log_audit(user, "Machine Status", row["machine_name"], status)
    return {"ok": True, "status": status}


def create_shift(body, admin):
    name = str(body.get("shift_name", "")).strip()
    if not name:
        raise InvalidRequest("shift_name is required")
    st, en = _to_datetime(body.get("start_time")), _to_datetime(body.get("end_time"))
    if not st or not en:
        raise InvalidRequest("start_time and end_time are required (HH:MM)")
    try:
        with get_cursor() as (_, cur):
            cur.execute("INSERT INTO shifts (shift_name, start_time, end_time) VALUES (%s,%s,%s)", (name, st.time(), en.time()))
    except IntegrityError as exc:
        raise Conflict("A shift with this name already exists") from exc
    log_audit(admin, "Created Shift", name, f"{st:%H:%M}-{en:%H:%M}")
    return {"ok": True}


def update_shift(shift_id, body, admin):
    fields = {}
    if "is_active" in body:
        fields["is_active"] = 1 if body["is_active"] else 0
    for key in ("start_time", "end_time"):
        if body.get(key):
            fields[key] = _to_datetime(body[key]).time()
    if body.get("shift_name"):
        fields["shift_name"] = str(body["shift_name"]).strip()
    if not fields:
        raise InvalidRequest("Nothing to update")
    with get_cursor() as (_, cur):
        cur.execute("SELECT shift_id FROM shifts WHERE shift_id = %s FOR UPDATE", (shift_id,))
        if not cur.fetchone():
            raise NotFound("Shift not found")
        if fields.get("is_active") == 0:
            cur.execute(
                "SELECT COUNT(*) AS n FROM users WHERE shift_id = %s AND role = 'SUPERVISOR'",
                (shift_id,),
            )
            if cur.fetchone()["n"]:
                raise Conflict("Reassign supervisors before deactivating this shift")
        cur.execute(f"UPDATE shifts SET {', '.join(k + ' = %s' for k in fields)} WHERE shift_id = %s", [*fields.values(), shift_id])
    log_audit(admin, "Updated Shift", f"#{shift_id}", ", ".join(fields))
    return {"ok": True}


def full_audit(limit=200):
    limit = max(1, min(int(limit), 1000))
    return [{"time": a["created_at"].strftime("%Y-%m-%d %I:%M:%S %p"), "user": a["user_name"], "action": a["action"],
             "target": a["target"], "details": a["details"]} for a in rows("SELECT * FROM audit_logs ORDER BY audit_id DESC LIMIT %s", (limit,))]


def change_password(user, current, new):
    from app.core.security import hash_password, verify_password
    if len(new or "") < 8:
        raise InvalidRequest("New password must contain at least 8 characters")
    row = one("SELECT password_hash FROM users WHERE user_id = %s", (user["user_id"],))
    if not row or not verify_password(current or "", row["password_hash"]):
        raise InvalidRequest("Current password is incorrect")
    with get_cursor() as (_, cur):
        cur.execute("UPDATE users SET password_hash = %s WHERE user_id = %s", (hash_password(new), user["user_id"]))
    log_audit(user, "Password Changed", user["email"])
    return {"ok": True}

def _totals(start, end):
    r = one(
        """SELECT COALESCE(SUM(quantity_produced),0) AS produced, COALESCE(SUM(quantity_defective),0) AS defective,
                  COALESCE(SUM(run_minutes),0) AS run, COALESCE(SUM(planned_minutes),0) AS planned
           FROM production_records WHERE production_date BETWEEN %s AND %s""", (start, end))
    t = one("SELECT COALESCE(SUM(target_quantity),0) AS t FROM production_targets WHERE target_date BETWEEN %s AND %s", (start, end))
    return r, float(t["t"])


def _production_report(start, end):
    production = rows(
        """SELECT production_date AS d,
                  SUM(quantity_produced) AS produced,
                  SUM(quantity_defective) AS defective
           FROM production_records
           WHERE production_date BETWEEN %s AND %s
           GROUP BY production_date ORDER BY production_date""",
        (start, end),
    )
    targets = rows(
        """SELECT target_date AS d, SUM(target_quantity) AS target
           FROM production_targets
           WHERE target_date BETWEEN %s AND %s
           GROUP BY target_date ORDER BY target_date""",
        (start, end),
    )
    product_rows = rows(
        """SELECT p.product_code, p.product_name,
                  COALESCE(r.produced, 0) AS produced,
                  COALESCE(r.defective, 0) AS defective,
                  COALESCE(t.target_quantity, 0) AS target
           FROM products p
           LEFT JOIN (
               SELECT product_id, SUM(quantity_produced) AS produced,
                      SUM(quantity_defective) AS defective
               FROM production_records
               WHERE production_date BETWEEN %s AND %s
               GROUP BY product_id
           ) r ON r.product_id = p.product_id
           LEFT JOIN (
               SELECT product_id, SUM(target_quantity) AS target_quantity
               FROM production_targets
               WHERE target_date BETWEEN %s AND %s
               GROUP BY product_id
           ) t ON t.product_id = p.product_id
           WHERE p.is_active = 1
           ORDER BY produced DESC""",
        (start, end, start, end),
    )
    production_by_day = {
        item["d"]: {
            "produced": int(item["produced"] or 0),
            "defective": int(item["defective"] or 0),
        }
        for item in production
    }
    targets_by_day = {item["d"]: int(item["target"] or 0) for item in targets}
    daily = []
    current = start
    while current <= end:
        totals = production_by_day.get(current, {})
        daily.append({
            "date": current.strftime("%d %b"),
            "produced": totals.get("produced", 0),
            "target": targets_by_day.get(current, 0),
            "defective": totals.get("defective", 0),
        })
        current += timedelta(days=1)
    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "produced": sum(item["produced"] for item in daily),
        "target": sum(item["target"] for item in daily),
        "defective": sum(item["defective"] for item in daily),
        "daily": daily,
        "products": [
            {
                "product_code": item["product_code"],
                "product_name": item["product_name"],
                "produced": int(item["produced"] or 0),
                "defective": int(item["defective"] or 0),
                "target": int(item["target"] or 0),
            }
            for item in product_rows
        ],
    }


def _oee(produced, defective, run, planned, target):
    """OEE = availability x performance x quality (percent values).
    Performance = output vs target; with no target set it is treated as 100%."""
    produced, defective, run, planned = map(float, (produced, defective, run, planned))
    availability = min(run / planned, 1) * 100 if planned else 0.0
    performance = min(produced / target, 1) * 100 if target else (100.0 if produced else 0.0)
    quality = (produced - defective) / produced * 100 if produced else 0.0
    return {"availability": round(availability, 1), "performance": round(performance, 1),
            "quality": round(quality, 1), "value": round(availability * performance * quality / 10000, 1)}


@ttl_cached(10)
def manager_dashboard(range_name="daily"):
    start, end, p_start, p_end = period(range_name)
    now = datetime.now()
    r, target = _totals(start, end)
    production_report = _production_report(start, end)
    cur_oee = _oee(r["produced"], r["defective"], r["run"], r["planned"], target)
    pr, pt = _totals(p_start, p_end)
    prev_oee = _oee(pr["produced"], pr["defective"], pr["run"], pr["planned"], pt)

    top = one(
        """SELECT p.product_code, p.product_name FROM production_records r JOIN products p ON p.product_id = r.product_id
           WHERE r.production_date BETWEEN %s AND %s GROUP BY p.product_id ORDER BY SUM(r.quantity_produced) DESC LIMIT 1""",
        (start, end))
    last7 = end - timedelta(days=6)
    day_totals = {x["d"]: float(x["q"]) for x in rows(
        "SELECT production_date AS d, SUM(quantity_produced) AS q FROM production_records WHERE production_date BETWEEN %s AND %s GROUP BY production_date",
        (last7, end))}
    bars = [round(day_totals.get(last7 + timedelta(days=i), 0) / 1000, 1) for i in range(7)]

    m = one("""SELECT COUNT(*) AS total, COALESCE(SUM(status='RUNNING'),0) AS running FROM machines""")
    lines = one("""SELECT COUNT(DISTINCT location) AS total,
                          COUNT(DISTINCT CASE WHEN status='RUNNING' THEN location END) AS active
                   FROM machines WHERE location IS NOT NULL AND location <> ''""")

    # 30-day OEE trend
    t_first = date.today() - timedelta(days=29)
    prod = {x["d"]: x for x in rows(
        """SELECT production_date AS d, SUM(quantity_produced) AS produced, SUM(quantity_defective) AS defective,
                  SUM(run_minutes) AS run, SUM(planned_minutes) AS planned
           FROM production_records WHERE production_date BETWEEN %s AND %s GROUP BY production_date""", (t_first, date.today()))}
    targ = {x["d"]: float(x["t"]) for x in rows(
        "SELECT target_date AS d, SUM(target_quantity) AS t FROM production_targets WHERE target_date BETWEEN %s AND %s GROUP BY target_date",
        (t_first, date.today()))}
    trend = []
    for i in range(30):
        d = t_first + timedelta(days=i)
        x = prod.get(d)
        o = _oee(x["produced"] or 0, x["defective"] or 0, x["run"] or 0, x["planned"] or 0, targ.get(d, 0)) if x else \
            {"availability": 0, "performance": 0, "quality": 0}
        trend.append({"date": d.strftime("%d %b"), "availability": o["availability"],
                      "performance": o["performance"], "quality": o["quality"]})

    # defect Pareto
    causes = rows(
        """SELECT dr.cause, SUM(dr.quantity) AS qty FROM defect_reasons dr
           JOIN production_records r ON r.production_id = dr.production_id
           WHERE r.production_date BETWEEN %s AND %s GROUP BY dr.cause ORDER BY qty DESC LIMIT 7""", (start, end))
    pareto = [{"cause": c["cause"], "count": int(c["qty"])} for c in causes]
    unexplained = int(r["defective"]) - sum(c["count"] for c in pareto)
    if unexplained > 0:
        pareto.append({"cause": "Unspecified", "count": unexplained})

    alerts = compute_alerts(start, end)

    plan_rows = rows(
        """SELECT s.shift_name, p.product_name, m.machine_name, m.location, t.target_quantity
           FROM production_targets t JOIN products p ON p.product_id = t.product_id
           LEFT JOIN shifts s ON s.shift_id = t.shift_id LEFT JOIN machines m ON m.machine_id = t.machine_id
           WHERE t.target_date = %s ORDER BY s.start_time""", (date.today(),))
    plan = [f"{(x['shift_name'] or 'Any shift')}: {x['product_name']} on {x['location'] or x['machine_name'] or 'any line'} (target {x['target_quantity']})"
            for x in plan_rows]

    return {
        "oee": {"value": cur_oee["value"], "delta": round(cur_oee["value"] - prev_oee["value"], 1),
                "availability": cur_oee["availability"], "performance": cur_oee["performance"], "quality": cur_oee["quality"]},
        "output": {"actual": int(r["produced"]), "plan": int(target),
                   "topProduct": f"{top['product_code']} {top['product_name']}" if top else "-", "bars": bars},
        "productionReport": production_report,
        "utilization": {"percent": pct(m["running"], m["total"]), "activeLines": int(lines["active"] or 0),
                        "totalLines": int(lines["total"] or 0), "activeMachines": int(m["running"]), "totalMachines": int(m["total"])},
        # No cost data is stored yet, so these stay at zero rather than showing made-up numbers.
        "cost": {"total": 0, "energy": 0, "labor": 0, "scrap": 0},
        "oeeTrend": trend,
        "pareto": pareto,
        "alerts": alerts,
        "machines": machine_floor(start, end),
        "plan": plan,
        "reports": REPORTS,
    }


# ---------------------------------------------------------------- admin
def list_manager_supervisors():
    return rows(
        """SELECT u.user_id, u.full_name, u.email, u.is_active,
                  s.shift_id, s.shift_name,
                  (SELECT r.production_date
                   FROM production_records r
                   WHERE r.supervisor_id = u.user_id
                   ORDER BY r.production_date DESC, r.recorded_at DESC
                   LIMIT 1) AS last_production_date,
                  (SELECT recorded_shift.shift_name
                   FROM production_records r
                   LEFT JOIN shifts recorded_shift ON recorded_shift.shift_id = r.shift_id
                   WHERE r.supervisor_id = u.user_id
                   ORDER BY r.production_date DESC, r.recorded_at DESC
                   LIMIT 1) AS last_production_shift
           FROM users u
           LEFT JOIN shifts s ON s.shift_id = u.shift_id
           WHERE u.role = 'SUPERVISOR'
           ORDER BY u.is_active DESC, u.full_name, u.user_id"""
    )


def assign_supervisor_shift(user_id, body, manager):
    shift_value = _pick(body, "shift_id", "shift")
    if shift_value is None:
        raise InvalidRequest("shift_id is required")
    shift = find_shift(shift_value)

    def validate_active_default_shift(selected):
        shift_name = str(selected["shift_name"]).strip()
        if not selected["is_active"]:
            raise InvalidRequest("The selected shift is inactive")
        shift_code = shift_name.split(maxsplit=1)[0].rstrip("-–—").upper() if shift_name else ""
        if shift_code not in {"A", "B", "C"}:
            raise InvalidRequest("Supervisors can only be assigned to A, B, or C shifts")

    validate_active_default_shift(shift)

    reason = str(body.get("reason") or "").strip()
    with get_cursor() as (_, cur):
        cur.execute(
            "SELECT shift_id, shift_name, is_active FROM shifts WHERE shift_id = %s FOR UPDATE",
            (shift["shift_id"],),
        )
        shift = cur.fetchone()
        if not shift:
            raise NotFound("Shift not found")
        validate_active_default_shift(shift)
        cur.execute(
            """SELECT user_id, full_name, email, role, is_active, shift_id
               FROM users WHERE user_id = %s FOR UPDATE""",
            (user_id,),
        )
        supervisor = cur.fetchone()
        if not supervisor or supervisor["role"] != "SUPERVISOR":
            raise NotFound("Supervisor not found")
        if not supervisor["is_active"]:
            raise InvalidRequest("Inactive supervisors cannot be assigned to a shift")

        previous_shift_id = supervisor["shift_id"]
        cur.execute(
            "UPDATE users SET shift_id = %s WHERE user_id = %s",
            (shift["shift_id"], user_id),
        )
        cur.execute(
            """SELECT u.user_id, u.full_name, u.email, u.is_active,
                      s.shift_id, s.shift_name
               FROM users u LEFT JOIN shifts s ON s.shift_id = u.shift_id
               WHERE u.user_id = %s""",
            (user_id,),
        )
        updated = cur.fetchone()

    change = (
        f"{supervisor['full_name']}: shift_id {previous_shift_id or 'unassigned'}"
        f" -> {shift['shift_id']}"
    )
    log_audit(
        manager,
        "Assigned Supervisor Shift",
        supervisor["email"],
        f"{change}; reason: {reason or 'not provided'}",
    )
    return updated


def unassign_supervisor_shift(user_id, body, manager):
    reason = str((body or {}).get("reason") or "").strip()
    with get_cursor() as (_, cur):
        cur.execute(
            """SELECT user_id, full_name, email, role, shift_id
               FROM users WHERE user_id = %s FOR UPDATE""",
            (user_id,),
        )
        supervisor = cur.fetchone()
        if not supervisor or supervisor["role"] != "SUPERVISOR":
            raise NotFound("Supervisor not found")
        if supervisor["shift_id"] is None:
            raise InvalidRequest("Supervisor does not have an assigned shift")

        previous_shift_id = supervisor["shift_id"]
        cur.execute("UPDATE users SET shift_id = NULL WHERE user_id = %s", (user_id,))
        cur.execute(
            """SELECT u.user_id, u.full_name, u.email, u.is_active,
                      s.shift_id, s.shift_name
               FROM users u LEFT JOIN shifts s ON s.shift_id = u.shift_id
               WHERE u.user_id = %s""",
            (user_id,),
        )
        updated = cur.fetchone()

    log_audit(
        manager,
        "Unassigned Supervisor Shift",
        supervisor["email"],
        f"{supervisor['full_name']}: shift_id {previous_shift_id} -> unassigned; "
        f"reason: {reason or 'not provided'}",
    )
    return updated


def admin_dashboard(range_name="today"):
    u = one("""SELECT COALESCE(SUM(is_active),0) AS active, COALESCE(SUM(role='SUPERVISOR'),0) AS operators,
                      COALESCE(SUM(role='MANAGER'),0) AS managers, COALESCE(SUM(role='ADMIN'),0) AS admins,
                      COALESCE(SUM(created_at >= %s),0) AS new FROM users""", (datetime.now() - timedelta(days=7),))
    p = one("""SELECT COUNT(*) AS total, COALESCE(SUM(is_active),0) AS active,
                      COALESCE(SUM(updated_at >= %s),0) AS today FROM products""", (datetime.combine(date.today(), dtime.min),))
    m = one("""SELECT COUNT(*) AS total, COALESCE(SUM(status='RUNNING'),0) AS running, COALESCE(SUM(status='MAINTENANCE'),0) AS maint,
                      COALESCE(SUM(status='IDLE'),0) AS idle, COALESCE(SUM(status='DOWN'),0) AS down FROM machines""")
    s = one("SELECT COUNT(*) AS total, COALESCE(SUM(is_active),0) AS active FROM shifts")
    total_cfg = int(p["total"]) + int(m["total"]) + int(s["total"])
    active_cfg = int(p["active"]) + int(m["total"] - m["down"]) + int(s["active"])
    last = one("SELECT MAX(created_at) AS t FROM audit_logs WHERE action NOT IN ('Login','Shift Entry','Downtime Logged','Downtime Ended')")
    ago = "never"
    if last and last["t"]:
        mins = int((datetime.now() - last["t"]).total_seconds() / 60)
        ago = fmt_minutes(mins) + " ago" if mins < 1440 else f"{mins // 1440}d ago"

    kpis = {
        "config": {"percent": pct(active_cfg, total_cfg), "active": active_cfg, "total": total_cfg, "lastChange": ago, "version": "0.1.0"},
        "users": {"active": int(u["active"]), "operators": int(u["operators"]), "managers": int(u["managers"]),
                  "admins": int(u["admins"]), "newRequests": int(u["new"])},
        "products": {"total": int(p["total"]), "active": int(p["active"]), "draft": 0,
                     "archived": int(p["total"] - p["active"]), "modifiedToday": int(p["today"])},
        # active = RUNNING, setup = IDLE, syncErrors = DOWN
        "machines": {"total": int(m["total"]), "active": int(m["running"]), "maintenance": int(m["maint"]),
                     "setup": int(m["idle"]), "syncErrors": int(m["down"])},
    }

    first = date.today() - timedelta(days=6)
    buckets = {first + timedelta(days=i): {"logins": 0, "dataEntry": 0, "adminActions": 0} for i in range(7)}
    for a in rows("SELECT DATE(created_at) AS d, action, COUNT(*) AS n FROM audit_logs WHERE DATE(created_at) >= %s GROUP BY DATE(created_at), action", (first,)):
        b = buckets.get(a["d"])
        if b is None:
            continue
        key = "logins" if a["action"] == "Login" else "dataEntry" if a["action"] in DATA_ENTRY_ACTIONS else "adminActions"
        b[key] += int(a["n"])
    activity = [{"day": DAY_NAMES[d.weekday()], **v} for d, v in buckets.items()]

    start, end, _, _ = period(range_name)
    recs = one("SELECT COUNT(*) AS n, COALESCE(SUM(shift_id IS NOT NULL),0) AS with_shift FROM production_records WHERE production_date BETWEEN %s AND %s", (start, end))
    prod_used = one("SELECT COUNT(DISTINCT product_id) AS n FROM production_records WHERE production_date BETWEEN %s AND %s", (start, end))
    mach_used = one("SELECT COUNT(DISTINCT machine_id) AS n FROM production_records WHERE production_date BETWEEN %s AND %s", (start, end))
    reconciliation = [
        {"list": "Products", "total": int(p["active"]), "pct": pct(prod_used["n"], p["active"]),
         "unreconciled": int(p["active"]) - int(prod_used["n"])},
        {"list": "Machines", "total": int(m["total"]), "pct": pct(mach_used["n"], m["total"]),
         "unreconciled": int(m["total"]) - int(mach_used["n"])},
        {"list": "Shifts", "total": int(recs["n"]), "pct": pct(recs["with_shift"], recs["n"]),
         "unreconciled": int(recs["n"]) - int(recs["with_shift"])},
    ]
    audit = [{"time": a["created_at"].strftime("%I:%M:%S %p"), "user": a["user_name"], "action": a["action"],
              "target": a["target"], "details": a["details"]}
             for a in rows("SELECT * FROM audit_logs ORDER BY audit_id DESC LIMIT 15")]
    return {"kpis": kpis, "activity": activity, "reconciliation": reconciliation, "audit": audit}


def list_users():
    return rows("SELECT user_id, full_name, email, role, is_active, created_at FROM users ORDER BY created_at DESC")


def update_user(user_id, body, admin):
    fields, labels = {}, []
    if "role" in body:
        role = str(body["role"]).upper()
        if role not in {"ADMIN", "SUPERVISOR", "MANAGER"}:
            raise InvalidRequest("Role must be ADMIN, SUPERVISOR or MANAGER")
        fields["role"] = role
        labels.append(f"role={role}")
    if "is_active" in body:
        fields["is_active"] = 1 if body["is_active"] else 0
        labels.append("active" if body["is_active"] else "deactivated")
    if "full_name" in body and str(body["full_name"]).strip():
        fields["full_name"] = str(body["full_name"]).strip()[:120]
        labels.append("name changed")
    if not fields:
        raise InvalidRequest("Nothing to update")
    if user_id == admin["user_id"] and (fields.get("is_active") == 0 or ("role" in fields and fields["role"] != "ADMIN")):
        raise InvalidRequest("You cannot remove your own admin access")
    with get_cursor() as (_, cur):
        cur.execute(f"UPDATE users SET {', '.join(k + ' = %s' for k in fields)} WHERE user_id = %s", [*fields.values(), user_id])
        cur.execute("SELECT user_id, full_name, email, role, is_active FROM users WHERE user_id = %s", (user_id,))
        row = cur.fetchone()
    if not row:
        raise NotFound("User not found")
    log_audit(admin, "Updated User", row["email"], ", ".join(labels))
    return row


def list_products(active_only=False):
    sql = "SELECT product_id, product_code, product_name, description, is_active FROM products"
    return rows(sql + (" WHERE is_active = 1" if active_only else "") + " ORDER BY product_code")


def create_product(body, admin):
    code, name = str(body.get("product_code", "")).strip(), str(body.get("product_name", "")).strip()
    if not code or not name:
        raise InvalidRequest("product_code and product_name are required")
    try:
        with get_cursor() as (_, cur):
            cur.execute("INSERT INTO products (product_code, product_name, description) VALUES (%s,%s,%s)", (code, name, body.get("description")))
            pid = cur.lastrowid
    except IntegrityError as exc:
        raise Conflict("Product code already exists") from exc
    log_audit(admin, "Created Product", f"{code} {name}")
    return one("SELECT product_id, product_code, product_name, description, is_active FROM products WHERE product_id = %s", (pid,))


def update_product(product_id, body, admin):
    fields = {k: body[k] for k in ("product_name", "description", "is_active") if k in body}
    if not fields:
        raise InvalidRequest("Nothing to update")
    if "is_active" in fields:
        fields["is_active"] = 1 if fields["is_active"] else 0
    with get_cursor() as (_, cur):
        cur.execute(f"UPDATE products SET {', '.join(k + ' = %s' for k in fields)} WHERE product_id = %s", [*fields.values(), product_id])
    row = one("SELECT product_id, product_code, product_name, description, is_active FROM products WHERE product_id = %s", (product_id,))
    if not row:
        raise NotFound("Product not found")
    log_audit(admin, "Updated Product", f"{row['product_code']} {row['product_name']}", ", ".join(fields))
    return row


def _hhmm(t):
    """MySQL TIME comes back as a timedelta; show it as HH:MM."""
    secs = int(t.total_seconds()) if hasattr(t, "total_seconds") else 0
    return f"{secs // 3600 % 24:02d}:{secs % 3600 // 60:02d}"


def list_shifts():
    ensure_default_shifts()
    return [{**r, "start_time": _hhmm(r["start_time"]), "end_time": _hhmm(r["end_time"])}
            for r in rows("SELECT shift_id, shift_name, start_time, end_time, is_active FROM shifts ORDER BY start_time")]


def create_target(body, user):
    product = find_product(_pick(body, "product", "product_id", "product_code"))
    machine = find_machine(body["machine"]) if _pick(body, "machine", "machine_id") else None
    shift = find_shift(_pick(body, "shift", "shift_id"))
    qty = _int(_pick(body, "quantity", "target_quantity", "target"), "Target quantity", 1)
    when = _pick(body, "date", "target_date")
    tdate = _to_datetime(f"{when} 00:00").date() if when else date.today()
    existing = one(
        """SELECT target_id FROM production_targets
           WHERE target_date = %s AND product_id = %s
             AND machine_id <=> %s AND shift_id <=> %s
           LIMIT 1""",
        (
            tdate,
            product["product_id"],
            machine["machine_id"] if machine else None,
            shift["shift_id"] if shift else None,
        ),
    )
    if existing:
        raise Conflict(
            "A target already exists for this product, date, machine, and shift. "
            "The saved target is fixed and was not changed."
        )
    with get_cursor() as (_, cur):
        cur.execute(
            "INSERT INTO production_targets (product_id, machine_id, shift_id, target_date, target_quantity, created_by) VALUES (%s,%s,%s,%s,%s,%s)",
            (product["product_id"], machine["machine_id"] if machine else None, shift["shift_id"] if shift else None, tdate, qty, user["user_id"]))
        tid = cur.lastrowid
    log_audit(user, "Set Target", f"{product['product_code']} on {tdate}", f"{qty} units")
    return {"target_id": tid, "ok": True}


def list_targets(day=None):
    d = _to_datetime(f"{day} 00:00").date() if day else date.today()
    return rows(
        """SELECT t.target_id, t.target_date, t.target_quantity, p.product_code, p.product_name,
                  m.machine_code, m.machine_name, s.shift_name,
                  COALESCE(
                      (SELECT SUM(r.quantity_produced)
                       FROM production_records r
                       WHERE r.production_date = t.target_date
                         AND r.product_id = t.product_id
                         AND (t.machine_id IS NULL OR r.machine_id = t.machine_id)
                         AND (t.shift_id IS NULL OR r.shift_id = t.shift_id)
                         AND (
                             t.shift_id IS NOT NULL
                             OR r.shift_id IS NULL
                             OR NOT EXISTS (
                                 SELECT 1
                                 FROM production_targets scoped
                                 WHERE scoped.target_date = t.target_date
                                   AND scoped.product_id = t.product_id
                                   AND scoped.shift_id = r.shift_id
                                   AND (
                                       scoped.machine_id IS NULL
                                       OR scoped.machine_id = r.machine_id
                                   )
                             )
                         )),
                      0
                  ) AS produced_quantity
           FROM production_targets t JOIN products p ON p.product_id = t.product_id
           LEFT JOIN machines m ON m.machine_id = t.machine_id LEFT JOIN shifts s ON s.shift_id = t.shift_id
           WHERE t.target_date = %s ORDER BY t.target_id DESC""", (d,))


def list_report_managers():
    return rows(
        """SELECT user_id, full_name
           FROM users
           WHERE role = 'MANAGER' AND is_active = 1
           ORDER BY full_name"""
    )


def create_target_progress_report(body, supervisor):
    target_id = _int(body.get("target_id"), "Target", minimum=1)
    manager_id = _int(body.get("manager_id"), "Manager", minimum=1)
    actual_quantity = _int(body.get("actual_quantity"), "Produced quantity")
    reason = str(body.get("reason") or "").strip()
    if len(reason) < 10:
        raise InvalidRequest("Please explain the issue in at least 10 characters.")
    if len(reason) > 2000:
        raise InvalidRequest("The explanation cannot exceed 2000 characters.")

    target = one(
        """SELECT t.target_id, t.target_quantity, p.product_code, p.product_name,
                  s.shift_name
           FROM production_targets t
           JOIN products p ON p.product_id = t.product_id
           LEFT JOIN shifts s ON s.shift_id = t.shift_id
           WHERE t.target_id = %s""",
        (target_id,),
    )
    if not target:
        raise NotFound("The selected production target was not found.")
    if actual_quantity >= int(target["target_quantity"]):
        raise InvalidRequest("This target is already complete; no incomplete-target report can be sent.")
    manager = one(
        "SELECT user_id FROM users WHERE user_id = %s AND role = 'MANAGER' AND is_active = 1",
        (manager_id,),
    )
    if not manager:
        raise NotFound("The selected active Manager was not found.")

    with get_cursor() as (_, cur):
        cur.execute(
            """INSERT INTO target_progress_reports
               (target_id, supervisor_id, manager_id, actual_quantity, reason)
               VALUES (%s, %s, %s, %s, %s)""",
            (
                target_id,
                supervisor["user_id"],
                manager_id,
                actual_quantity,
                reason,
            ),
        )
        report_id = cur.lastrowid
    return {"report_id": report_id, "manager_id": manager_id}


def list_manager_target_reports(manager):
    return rows(
        """SELECT r.report_id, r.actual_quantity, r.reason, r.created_at,
                  r.reviewed_at, t.target_date, t.target_quantity,
                  p.product_code, p.product_name, s.shift_name,
                  sender.full_name AS supervisor_name
           FROM target_progress_reports r
           JOIN production_targets t ON t.target_id = r.target_id
           JOIN products p ON p.product_id = t.product_id
           LEFT JOIN shifts s ON s.shift_id = t.shift_id
           JOIN users sender ON sender.user_id = r.supervisor_id
           WHERE r.manager_id = %s
           ORDER BY r.created_at DESC
           LIMIT 100""",
        (manager["user_id"],),
    )


def review_target_progress_report(report_id, manager):
    with get_cursor() as (_, cur):
        cur.execute(
            """UPDATE target_progress_reports
               SET reviewed_at = CURRENT_TIMESTAMP
               WHERE report_id = %s AND manager_id = %s AND reviewed_at IS NULL""",
            (report_id, manager["user_id"]),
        )
        if cur.rowcount == 0:
            cur.execute(
                "SELECT report_id FROM target_progress_reports WHERE report_id = %s AND manager_id = %s",
                (report_id, manager["user_id"]),
            )
            existing = cur.fetchone()
            if not existing:
                raise NotFound("Target progress report was not found.")
    return {"report_id": report_id, "reviewed": True}


def list_product_production(day=None, shift_name=None):
    production_date = _to_datetime(f"{day} 00:00").date() if day else date.today()
    shift_filter = shift_name or None
    return rows(
        """SELECT p.product_id, p.product_code, p.product_name,
                  COALESCE(t.target_quantity, 0) AS target_quantity,
                  COALESCE(r.produced_quantity, 0) AS produced_quantity
           FROM products p
           LEFT JOIN (
               SELECT product_id, SUM(target_quantity) AS target_quantity
               FROM production_targets
               WHERE target_date = %s
                 AND (%s IS NULL OR shift_id IS NULL OR shift_id = (
                     SELECT shift_id FROM shifts WHERE LEFT(shift_name, 1) = %s LIMIT 1
                 ))
               GROUP BY product_id
           ) t ON t.product_id = p.product_id
           LEFT JOIN (
               SELECT product_id, SUM(quantity_produced) AS produced_quantity
               FROM production_records
               WHERE production_date = %s
                 AND (%s IS NULL OR shift_id IS NULL OR shift_id = (
                     SELECT shift_id FROM shifts WHERE LEFT(shift_name, 1) = %s LIMIT 1
                 ))
               GROUP BY product_id
           ) r ON r.product_id = p.product_id
           WHERE p.is_active = 1
           ORDER BY p.product_code""",
        (
            production_date,
            shift_filter,
            shift_filter,
            production_date,
            shift_filter,
            shift_filter,
        ),
    )


# ---------------------------------------------------------------- reports
REPORTS = [
    {"name": "Full Daily Summary", "type": "PDF"}, {"name": "Weekly OEE Report", "type": "Excel"},
    {"name": "Downtime Log", "type": "PDF"}, {"name": "Scrap & Defect Analysis", "type": "PDF"},
    {"name": "Machine Performance", "type": "Excel"},
]


def build_report(name):
    """Returns the column headers and rows used to render the requested report."""
    today, month_ago, week_ago = date.today(), date.today() - timedelta(days=29), date.today() - timedelta(days=6)
    if name == "Full Daily Summary":
        data = rows("""SELECT p.product_code, p.product_name, m.machine_name, COALESCE(s.shift_name,'-') AS shift_name,
                              r.quantity_produced, r.quantity_defective, r.run_minutes, r.planned_minutes
                       FROM production_records r JOIN products p ON p.product_id=r.product_id JOIN machines m ON m.machine_id=r.machine_id
                       LEFT JOIN shifts s ON s.shift_id=r.shift_id WHERE r.production_date=%s ORDER BY r.production_id""", (today,))
        return ["Product code", "Product", "Machine", "Shift", "Produced", "Defective", "Run min", "Planned min"], [list(r.values()) for r in data]
    if name == "Weekly OEE Report":
        out = []
        prod = {x["d"]: x for x in rows("""SELECT production_date AS d, SUM(quantity_produced) AS produced, SUM(quantity_defective) AS defective,
                SUM(run_minutes) AS run, SUM(planned_minutes) AS planned FROM production_records WHERE production_date BETWEEN %s AND %s GROUP BY production_date""", (week_ago, today))}
        targ = {x["d"]: float(x["t"]) for x in rows("SELECT target_date AS d, SUM(target_quantity) AS t FROM production_targets WHERE target_date BETWEEN %s AND %s GROUP BY target_date", (week_ago, today))}
        for i in range(7):
            d = week_ago + timedelta(days=i)
            x = prod.get(d)
            o = _oee(x["produced"] or 0, x["defective"] or 0, x["run"] or 0, x["planned"] or 0, targ.get(d, 0)) if x else {"availability": 0, "performance": 0, "quality": 0, "value": 0}
            out.append([d.isoformat(), o["availability"], o["performance"], o["quality"], o["value"]])
        return ["Date", "Availability %", "Performance %", "Quality %", "OEE %"], out
    if name == "Downtime Log":
        data = rows("""SELECT m.machine_name, d.reason, d.started_at, d.ended_at,
                              TIMESTAMPDIFF(MINUTE, d.started_at, COALESCE(d.ended_at, %s)) AS minutes
                       FROM downtime_events d JOIN machines m ON m.machine_id=d.machine_id
                       WHERE DATE(d.started_at) BETWEEN %s AND %s ORDER BY d.started_at DESC""", (datetime.now(), month_ago, today))
        return ["Machine", "Reason", "Started", "Ended", "Minutes"], [[r["machine_name"], r["reason"], r["started_at"], r["ended_at"] or "Ongoing", r["minutes"]] for r in data]
    if name == "Scrap & Defect Analysis":
        data = rows("""SELECT m.machine_name, SUM(r.quantity_produced) AS produced, SUM(r.quantity_defective) AS defective
                       FROM production_records r JOIN machines m ON m.machine_id=r.machine_id
                       WHERE r.production_date BETWEEN %s AND %s GROUP BY m.machine_id ORDER BY defective DESC""", (month_ago, today))
        return ["Machine", "Produced", "Defective", "Defect rate %"], [[r["machine_name"], r["produced"], r["defective"], pct(r["defective"], r["produced"])] for r in data]
    if name == "Machine Performance":
        data = rows("""SELECT m.machine_name, m.status, COALESCE(SUM(r.quantity_produced),0) AS produced, COALESCE(SUM(r.quantity_defective),0) AS defective,
                              COALESCE(SUM(r.run_minutes),0) AS run, COALESCE(SUM(r.planned_minutes),0) AS planned
                       FROM machines m LEFT JOIN production_records r ON r.machine_id=m.machine_id AND r.production_date BETWEEN %s AND %s
                       GROUP BY m.machine_id ORDER BY m.machine_name""", (month_ago, today))
        return ["Machine", "Status", "Produced", "Defective", "Run min", "Planned min", "Availability %"], \
               [[r["machine_name"], r["status"], r["produced"], r["defective"], r["run"], r["planned"], pct(r["run"], r["planned"])] for r in data]
    raise NotFound(f"'{name}' is not available yet: no data is tracked for it")


def report_csv(name):
    header, data = build_report(name)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(data)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def report_excel(name):
    header, data = build_report(name)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\[\]:*?/\\]", " ", name)[:31] or "Report"
    sheet.append(header)
    for row in data:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="0D5A46")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column]
        width = min(max(max(values, default=0) + 2, 12), 40)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_pdf(name):
    header, data = build_report(name)
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=name,
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 8
    cell_style.leading = 10
    table_data = [
        [Paragraph(escape(str(value if value is not None else "")), cell_style) for value in header]
    ]
    if data:
        table_data.extend([
            [Paragraph(escape(str(value if value is not None else "")), cell_style) for value in row]
            for row in data
        ])
    else:
        table_data.append([
            Paragraph("No records for this report period.", cell_style),
            *[Paragraph("", cell_style) for _ in header[1:]],
        ])
    report_table = Table(table_data, repeatRows=1, hAlign="LEFT")
    report_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D5A46")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D5E3DE")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F7F4")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    document.build([Paragraph(escape(name), styles["Title"]), Spacer(1, 5 * mm), report_table])
    return output.getvalue()


def report_file(name, file_type):
    normalized_type = str(file_type or "").strip().lower()
    if normalized_type == "pdf":
        return report_pdf(name), "application/pdf", "pdf"
    if normalized_type in {"excel", "xlsx"}:
        return (
            report_excel(name),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    raise InvalidRequest("Report format must be PDF or Excel")